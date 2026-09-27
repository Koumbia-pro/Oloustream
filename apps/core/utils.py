"""Petits utilitaires partagés entre les applications."""
import logging

from django.http import HttpResponse
from django.shortcuts import redirect
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme

logger = logging.getLogger(__name__)


def safe_redirect(request, url, fallback):
    """Redirige vers `url` uniquement si elle pointe vers ce site, sinon vers `fallback`."""
    if url and url_has_allowed_host_and_scheme(
        url, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        return redirect(url)
    return redirect(fallback)


def parse_int(value, default=None):
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def querystring_without_page(request):
    params = request.GET.copy()
    params.pop("page", None)
    return params.urlencode()


def send_safely(send_callable, *args, **kwargs):
    """
    Exécute un envoi d'email sans jamais faire planter la requête.
    Retourne True si l'envoi a réussi.
    """
    try:
        send_callable(*args, **kwargs)
        return True
    except Exception:  # SMTP indisponible, mauvais identifiants, etc.
        logger.exception("Échec d'envoi d'email via %s", getattr(send_callable, "__name__", send_callable))
        return False


# ---------- Export Excel ----------

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _excel_safe(value):
    """Empêche l'injection de formules dans Excel (valeurs saisies par les utilisateurs)."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def excel_response(filename_prefix, sheet_title, headers, rows, column_widths=None):
    """Construit une réponse HTTP contenant un fichier XLSX mis en forme."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    ws.append(headers)
    for row in rows:
        ws.append([_excel_safe(v) for v in row])

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="0F172A")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

    for index, width in enumerate(column_widths or [], start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    timestamp = timezone.localtime().strftime("%Y%m%d_%H%M")
    response["Content-Disposition"] = f'attachment; filename="{filename_prefix}_oloustream_{timestamp}.xlsx"'
    wb.save(response)
    return response


def fmt_date(value, pattern="%d/%m/%Y"):
    if not value:
        return ""
    if hasattr(value, "tzinfo") and value.tzinfo is not None:
        value = timezone.localtime(value)
    return value.strftime(pattern)


def fmt_datetime(value):
    return fmt_date(value, "%d/%m/%Y %H:%M")


def yes_no(value):
    return "Oui" if value else "Non"
